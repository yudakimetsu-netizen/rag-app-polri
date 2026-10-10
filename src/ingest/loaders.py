import os
from pathlib import Path
from langchain_core.documents import Document

def load_pdf(file_path: Path) -> list[Document]:
    """Load PDF menggunakan PyPDFLoader."""
    from langchain_community.document_loaders import PyPDFLoader
    loader = PyPDFLoader(str(file_path))
    return loader.load()

def load_docx(file_path: Path) -> list[Document]:
    """Load DOCX menggunakan python-docx."""
    from docx import Document as DocxDocument
    doc = DocxDocument(str(file_path))
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)
    content = "\n".join(full_text)
    return [Document(page_content=content, metadata={"page": 1})]

def load_txt(file_path: Path) -> list[Document]:
    """Load TXT/MD sederhana."""
    with open(file_path, "r", encoding="utf-8") as f:
        content = f.read()
    return [Document(page_content=content, metadata={"page": 1})]

# Mapping ekstensi ke loader
LOADERS = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".txt": load_txt,
    ".md": load_txt,
}

def load_document(file_path: Path) -> list[Document]:
    """Load satu file berdasarkan ekstensinya."""
    ext = file_path.suffix.lower()
    if ext not in LOADERS:
        print(f"  [SKIP] Format tidak didukung: {ext}")
        return []
    
    try:
        docs = LOADERS[ext](file_path)
        print(f"  [OK]   {file_path.name} ({len(docs)} halaman)")
        return docs
    except Exception as e:
        print(f"  [FAIL] {file_path.name}: {e}")
        return []

def load_all_documents(knowledge_dir: Path) -> list[Document]:
    """Load semua file dari folder knowledge."""
    documents = []
    print(f"Memuat dokumen dari: {knowledge_dir}")
    
    if not knowledge_dir.exists():
        print(f"  [ERROR] Folder tidak ditemukan: {knowledge_dir}")
        return []
    
    files = sorted(knowledge_dir.iterdir())
    for file_path in files:
        if file_path.is_file() and not file_path.name.startswith("."):
            docs = load_document(file_path)
            # Tambahkan metadata sumber ke setiap halaman
            for doc in docs:
                doc.metadata["source"] = file_path.name
                doc.metadata["source_path"] = str(file_path)
            documents.extend(docs)
    
    print(f"Total: {len(documents)} halaman dari {len(files)} file")
    return documents


if __name__ == "__main__":
    from src.config import KNOWLEDGE_DIR
    docs = load_all_documents(KNOWLEDGE_DIR)
    if docs:
        print("\nContoh metadata halaman pertama:")
        print(docs[0].metadata)